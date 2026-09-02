import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_attendance_excel(records, institution_name: str = "Apex Institute of Technology") -> io.BytesIO:
    """
    Generates a beautifully styled Excel (.xlsx) attendance report workbook with summary block.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    # Color definitions
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    present_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    # Title Block
    ws['A1'] = institution_name.upper()
    ws['A1'].font = title_font
    ws['A2'] = f"OFFICIAL ATTENDANCE REPORT — Generated on {datetime.now().strftime('%d %B %Y, %I:%M %p')}"
    ws['A2'].font = subtitle_font

    # Summary Row Stats
    total_records = len(records)
    present_count = sum(1 for r in records if r.status in ('Present', 'Late', 'Half Day'))
    absent_count = sum(1 for r in records if r.status == 'Absent')
    rate = round((present_count / total_records * 100), 1) if total_records > 0 else 0.0

    ws['A4'] = "Summary Metrics:"
    ws['A4'].font = bold_font
    ws['B4'] = f"Total Records: {total_records}"
    ws['C4'] = f"Present: {present_count}"
    ws['D4'] = f"Absent: {absent_count}"
    ws['E4'] = f"Attendance Rate: {rate}%"
    for col in ['B4', 'C4', 'D4', 'E4']:
        ws[col].font = bold_font

    # Table Headers at Row 6
    headers = [
        "S.No", "Student ID", "Student Name", "Department",
        "Date", "Time In", "Time Out", "Status", "Method", "Marked By"
    ]
    header_row = 6
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 24

    # Data Rows
    current_row = 7
    for idx, r in enumerate(records, 1):
        s = r.student
        dept_code = s.department.code if (s and s.department) else "N/A"
        date_str = r.date.strftime('%Y-%m-%d') if r.date else ""
        in_str = r.time_in.strftime('%I:%M %p') if r.time_in else "-"
        out_str = r.time_out.strftime('%I:%M %p') if r.time_out else "-"
        marker_name = r.marker.username if r.marker else "System"

        row_data = [
            idx,
            s.student_id if s else "N/A",
            s.full_name if s else "N/A",
            dept_code,
            date_str,
            in_str,
            out_str,
            r.status,
            r.method,
            marker_name
        ]

        is_even = (idx % 2 == 0)
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = normal_font
            cell.border = thin_border

            # Specific column alignments & fills
            if col_idx in (1, 2, 5, 6, 7, 9):
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")

            if col_idx == 8: # Status column
                if r.status == 'Present':
                    cell.fill = present_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="15803D")
                elif r.status == 'Absent':
                    cell.fill = absent_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="B91C1C")
                cell.alignment = Alignment(horizontal="center")
            elif is_even:
                cell.fill = zebra_fill

        ws.row_dimensions[current_row].height = 18
        current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 6:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
